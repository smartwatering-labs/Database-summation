from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json, hashlib, zipfile

WORK = Path(__file__).resolve().parent
OUT = WORK.parent

def font(style, size, bold=False):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    rf.set(qn('w:eastAsia'), '微软雅黑')
    rf.set(qn('w:ascii'), 'Arial'); rf.set(qn('w:hAnsi'), 'Arial')
    for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        rf.attrib.pop(qn('w:'+attr),None)

def cell_content(cell, value, header=False):
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.line_spacing = Pt(14)
        for r in p.runs:
            r.font.size = Pt(9.5)
            r.font.bold = header

def table(doc, rows):
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [5.1, 11.4] if len(rows[0]) == 2 else [16.5 / len(rows[0])] * len(rows[0])
    for c, width in zip(t.columns, widths): c.width = Cm(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    props.append(borders)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (cell, value) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[j]); cell_content(cell, value, i == 0)
            pr = cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                e=OxmlElement('w:'+side); e.set(qn('w:w'),'75'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            if i == 0:
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'E7EEF4'); pr.append(shd)
        trpr=t.rows[i]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit'); trpr.append(no_split)
        if i == 0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

manifest = []
for name in ('硬件1','硬件2','网页1','网页2'):
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8)
    sec.left_margin=Cm(2.1); sec.right_margin=Cm(2.1)
    grid=sec._sectPr.find(qn('w:docGrid'))
    if grid is not None: sec._sectPr.remove(grid)
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'): border.getparent().remove(border)
    font(doc.styles['Normal'],10.5)
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=Pt(16); normal.space_after=Pt(6)
    normal.widow_control=True
    snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0')
    doc.styles['Normal'].element.get_or_add_pPr().append(snap)
    font(doc.styles['Title'],20,True)
    doc.styles['Title'].paragraph_format.space_after=Pt(10)
    font(doc.styles['Subtitle'],10)
    doc.styles['Subtitle'].paragraph_format.space_after=Pt(10)
    font(doc.styles['Heading 1'],14,True)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(8)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
    doc.styles['Heading 1'].paragraph_format.keep_with_next=True
    font(doc.styles['Heading 2'],12,True)
    for st in ('List Bullet',):
        font(doc.styles[st],9.5)
        doc.styles[st].paragraph_format.space_after=Pt(4)
        doc.styles[st].paragraph_format.line_spacing=Pt(14)
    # Page numbers help later task discussions without decorative headers.
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run(); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); r._r.addnext(field)
    lines=(WORK/(name+'.md')).read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line=='<!-- page -->': doc.add_page_break()
        elif line.startswith('# '): doc.add_paragraph(line[2:],'Title')
        elif line.startswith('## '): doc.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('修订日期 '): doc.add_paragraph(line,'Subtitle')
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(set(x)<=set('-: ') for x in cells): rows.append(cells)
                i+=1
            table(doc,rows); continue
        elif line.startswith('- '): doc.add_paragraph(line[2:],'List Bullet')
        else: doc.add_paragraph(line)
        i+=1
    doc.core_properties.title=lines[0][2:]
    doc.core_properties.subject='依据浇水平台现有代码与10-01硬件资料修订的改进方案'
    doc.core_properties.author=''
    doc.core_properties.keywords='项目背景修订 方案 待实现 真机待验证'
    dest=OUT/(name+' 项目背景修订版.docx'); doc.save(dest)
    with zipfile.ZipFile(dest) as z:
        assert z.testzip() is None
    assert not any('<!--' in p.text for p in doc.paragraphs)
    manifest.append({'name':name,'output':str(dest),'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'plannedPages':1+lines.count('<!-- page -->'),'characters':sum(len(p.text) for p in doc.paragraphs)})
    print(dest.name, 'planned pages', manifest[-1]['plannedPages'])
    original=Path('C:/Users/bYYYyyL/Desktop')/(name+'.docx')
    with zipfile.ZipFile(original) as z:
        from lxml import etree
        tree=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        paragraphs=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in tree.xpath('//w:body//w:p',namespaces=ns)]
    (WORK/(name+' 原文提取.txt')).write_text('\n'.join(paragraphs),encoding='utf-8')
    manifest[-1]['originalSha256']=hashlib.sha256(original.read_bytes()).hexdigest()
(WORK/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
