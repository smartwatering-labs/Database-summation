from pathlib import Path
import sys
sys.stdout.reconfigure(encoding='utf-8')
root=Path(__file__).resolve().parent
source=(root/'build_documents.py').read_text(encoding='utf-8')
prefix=source.split('manifest = []')[0]
prefix+='''
def linked_paragraph(doc, line):
    import re
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    p=doc.add_paragraph()
    cursor=0
    for match in re.finditer(r'\\[([^\\]]+)\\]\\((https?://[^)]+)\\)',line):
        p.add_run(line[cursor:match.start()])
        link=OxmlElement('w:hyperlink')
        link.set(qn('r:id'),p.part.relate_to(match.group(2),RT.HYPERLINK,is_external=True))
        run=OxmlElement('w:r')
        t=OxmlElement('w:t'); t.text=match.group(1); run.append(t)
        link.append(run); p._p.append(link)
        cursor=match.end()
    p.add_run(line[cursor:])
'''
body=source.split("for name in ('硬件1','硬件2','网页1','网页2'):\n",1)[1].split('    manifest.append(',1)[0]
body=body.replace("(WORK/(name+'.md'))", "(OUT/(name+'.md'))")
body=body.replace("(name+' 项目背景修订版.docx')", "(name+'.docx')")
body=body.replace("else: doc.add_paragraph(line)", "else: linked_paragraph(doc,line)")
body=body.replace("widths = [5.1, 11.4]", "widths = [5.1, 11.4]")
code=prefix+"\nfor name in ('可行性复核与实施计划',):\n"+body+"\n    print(dest)\n"
exec(compile(code, str(root/'build_plan.py'), 'exec'))
