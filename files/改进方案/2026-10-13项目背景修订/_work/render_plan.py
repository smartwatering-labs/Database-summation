from pathlib import Path
import sys, json, re
sys.stdout.reconfigure(encoding='utf-8')
import pypdfium2 as pdfium
import pdfplumber
from docx import Document
root=Path(__file__).resolve().parent
out=root/'render'/'计划'
pdf=pdfium.PdfDocument(str(out/'计划.pdf'))
for index in range(len(pdf)):
    page=pdf[index]
    page.render(scale=1.5).to_pil().save(out/f'page-{index+1}.png')
    page.close()
pdf.close()
details=[]
with pdfplumber.open(out/'计划.pdf') as pdf:
    full=[]
    for index, page in enumerate(pdf.pages):
        text=page.extract_text() or ''
        full.append(text)
        for table in page.find_tables():
            for row in table.extract():
                full.extend(value for value in row if value)
        lines=text.splitlines()
        details.append({'page':index+1, 'characters':len(text),'first':lines[:3],'last':lines[-3:]})
        print(json.dumps(details[-1],ensure_ascii=False))
    rendered=re.sub(r'\s+','', '\n'.join(full))
doc=Document(root.parent/'可行性复核与实施计划.docx')
paras=[p.text for p in doc.paragraphs if p.text.strip()]
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            paras.extend(p.text for p in cell.paragraphs if p.text.strip())
missing=[p for p in paras if re.sub(r'\s+','',p) not in rendered]
print('Missing paragraphs',len(missing))
if missing: print(missing)
(out/'verification.json').write_text(json.dumps({'pages':details,'missing':missing},ensure_ascii=False,indent=2),encoding='utf-8')
